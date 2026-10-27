import io
import fitz
import zipfile


def upload_pdf(client, path):
    with path.open("rb") as handle:
        response = client.post("/api/files", files={"file": (path.name, handle, "application/pdf")})
    assert response.status_code == 200, response.text
    return response.json()["data"]


def test_health_reports_optional_capabilities(client):
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json()["capabilities"]["pdf"] is True
    assert "ocr" in response.json()["capabilities"]
    assert response.headers["X-Content-Type-Options"] == "nosniff"
    assert response.headers["Cache-Control"] == "no-store"
    assert response.headers["Referrer-Policy"] == "no-referrer"


def test_document_identifier_cannot_escape_workspace(client):
    response = client.get("/api/files/..%2F..%2FWindows/download")
    assert response.status_code in {404, 422}


def test_local_api_token_protects_everything_except_health(client, monkeypatch):
    import backend.app.main as main_module
    monkeypatch.setattr(main_module, "API_TOKEN", "test-session-token")
    assert client.get("/api/health").status_code == 200
    assert client.get("/api/docs").status_code == 401
    assert client.get("/api/docs", headers={"X-Atelier-Token": "wrong"}).status_code == 401
    assert client.get("/api/docs", headers={"X-Atelier-Token": "test-session-token"}).status_code == 200


def test_upload_render_and_extract_text(client, sample_pdf):
    uploaded = upload_pdf(client, sample_pdf)
    assert uploaded["kind"] == "pdf"
    assert uploaded["pdf"]["page_count"] == 1

    rendered = client.get(f"/api/pdf/{uploaded['id']}/pages/0/render")
    assert rendered.status_code == 200
    assert rendered.headers["content-type"] == "image/png"
    assert rendered.content.startswith(b"\x89PNG")

    blocks = client.get(f"/api/pdf/{uploaded['id']}/pages/0/text").json()["data"]
    assert any("Eski metin 12345" in block["text"] for block in blocks)


def test_edit_removes_old_text_from_content_and_persists(client, sample_pdf):
    uploaded = upload_pdf(client, sample_pdf)
    blocks = client.get(f"/api/pdf/{uploaded['id']}/pages/0/text").json()["data"]
    block = next(item for item in blocks if "Eski metin 12345" in item["text"])
    x0, y0, x1, y1 = block["bbox"]
    response = client.post(f"/api/pdf/{uploaded['id']}/edit-text", json={
        "page": 0,
        "rect": [x0, y0, max(x1 + 80, x0 + 250), y1 + 12],
        "text": "Yeni metin Türkçe 67890",
        "original_text": "Eski metin 12345",
        "style": {"font_size": 11, "color": "#172033", "align": "left"},
    })
    assert response.status_code == 200, response.text

    downloaded = client.get(f"/api/files/{uploaded['id']}/download")
    assert downloaded.status_code == 200
    with fitz.open(stream=downloaded.content, filetype="pdf") as output:
        text = "\n".join(page.get_text() for page in output).replace("\u00a0", " ")
        assert "Eski metin 12345" not in text
        assert "Yeni metin" in text
        assert "Türkçe" in text


def test_text_block_can_be_drag_moved_with_original_box_size(client, sample_pdf):
    uploaded = upload_pdf(client, sample_pdf)
    document_id = uploaded["id"]
    block = next(item for item in client.get(f"/api/pdf/{document_id}/pages/0/text").json()["data"] if "Eski metin 12345" in item["text"])
    x0, y0, x1, y1 = block["bbox"]
    target = [x0 + 45, y0 + 90, x1 + 45, y1 + 90]
    moved = client.post(f"/api/pdf/{document_id}/edit-text", json={
        "page": 0,
        "rect": block["bbox"],
        "target_rect": target,
        "text": block["text"],
        "original_text": block["text"],
        "style": {"font_size": block["font_size"], "font_family": block["font"], "color": block["color"], "align": "left"},
    })
    assert moved.status_code == 200, moved.text
    content = client.get(f"/api/files/{document_id}/download").content
    with fitz.open(stream=content, filetype="pdf") as output:
        matches = output[0].search_for("Eski metin 12345")
        assert len(matches) == 1
        assert abs(matches[0].x0 - target[0]) < 3
        assert abs(matches[0].y0 - target[1]) < 5
        assert not output[0].get_textbox(fitz.Rect(block["bbox"])).strip()


def test_table_cells_are_independent_and_edit_keeps_neighbor_positions(client, table_layout_pdf):
    uploaded = upload_pdf(client, table_layout_pdf)
    document_id = uploaded["id"]
    blocks = client.get(f"/api/pdf/{document_id}/pages/0/text").json()["data"]

    # Distant cells and adjacent visual lines must never become one editable box.
    assert next(item for item in blocks if item["text"] == "TOPLAM (TOTAL) (TL)")
    amount = next(item for item in blocks if item["text"] == "3.142,15")
    assert next(item for item in blocks if item["text"] == "8.142,15")
    assert next(item for item in blocks if item["text"] == "ALT SATIR")

    with fitz.open(table_layout_pdf) as original:
        label_before = original[0].search_for("TOPLAM (TOTAL) (TL)")[0]
        lower_before = original[0].search_for("ALT SATIR")[0]

    edited = client.post(f"/api/pdf/{document_id}/edit-text", json={
        "page": 0,
        "rect": amount["bbox"],
        "target_rect": amount["bbox"],
        "source_baseline": amount["baseline"],
        "text": "9.876,54",
        "original_text": amount["text"],
        "style": {
            "font_size": amount["font_size"], "font_family": amount["font"],
            "color": amount["color"], "align": "left",
        },
    })
    assert edited.status_code == 200, edited.text

    content = client.get(f"/api/files/{document_id}/download").content
    with fitz.open(stream=content, filetype="pdf") as output:
        page = output[0]
        label_after = page.search_for("TOPLAM (TOTAL) (TL)")[0]
        lower_after = page.search_for("ALT SATIR")[0]
        new_amount = page.search_for("9.876,54")[0]
        assert abs(label_after.x0 - label_before.x0) < 0.5
        assert abs(label_after.y0 - label_before.y0) < 0.5
        assert abs(lower_after.x0 - lower_before.x0) < 0.5
        assert abs(lower_after.y0 - lower_before.y0) < 0.5
        assert abs(new_amount.x0 - amount["bbox"][0]) < 1
        assert abs(new_amount.y0 - amount["bbox"][1]) < 1


def test_search_and_replace_removes_all_matching_text_and_supports_undo(client, sample_pdf):
    uploaded = upload_pdf(client, sample_pdf)
    document_id = uploaded["id"]

    found = client.get(f"/api/pdf/{document_id}/search", params={"query": "metin"})
    assert found.status_code == 200, found.text
    assert found.json()["data"]["count"] == 1
    assert found.json()["data"]["matches"][0]["page"] == 0

    replaced = client.post(f"/api/pdf/{document_id}/replace-text", json={
        "query": "metin",
        "replacement": "icerik",
        "case_sensitive": False,
        "whole_word": True,
    })
    assert replaced.status_code == 200, replaced.text
    assert replaced.json()["data"]["audit"]["matches"] == 1
    downloaded = client.get(f"/api/files/{document_id}/download")
    with fitz.open(stream=downloaded.content, filetype="pdf") as output:
        text = "\n".join(page.get_text() for page in output)
        assert "Eski metin 12345" not in text
        assert "icerik" in text

    undone = client.post(f"/api/pdf/{document_id}/undo")
    assert undone.status_code == 200, undone.text
    downloaded = client.get(f"/api/files/{document_id}/download")
    with fitz.open(stream=downloaded.content, filetype="pdf") as output:
        assert "Eski metin 12345" in "\n".join(page.get_text() for page in output)


def test_text_block_can_move_resize_and_apply_bold_italic_style(client, sample_pdf):
    uploaded = upload_pdf(client, sample_pdf)
    document_id = uploaded["id"]
    blocks = client.get(f"/api/pdf/{document_id}/pages/0/text").json()["data"]
    source = next(item for item in blocks if "Ikinci satir" in item["text"])
    response = client.post(f"/api/pdf/{document_id}/edit-text", json={
        "page": 0,
        "rect": source["bbox"],
        "target_rect": [300, 120, 520, 175],
        "text": "Tasinan stilli metin",
        "original_text": source["text"],
        "style": {
            "font_size": 13, "color": "#123456", "align": "center",
            "bold": True, "italic": True, "underline": True,
        },
    })
    assert response.status_code == 200, response.text
    updated = client.get(f"/api/pdf/{document_id}/pages/0/text").json()["data"]
    moved = next(item for item in updated if "Tasinan stilli metin" in item["text"].replace("\u00a0", " "))
    assert moved["bbox"][0] >= 300
    assert moved["bold"] is True
    assert moved["italic"] is True
    assert all("Ikinci satir" not in item["text"] for item in updated)


def test_split_and_rotate(client, sample_pdf):
    uploaded = upload_pdf(client, sample_pdf)
    rotated = client.post(f"/api/pdf/{uploaded['id']}/rotate", json={"pages": [0], "angle": 90})
    assert rotated.status_code == 200
    assert rotated.json()["data"]["pages"][0]["rotation"] == 90
    split = client.post(f"/api/pdf/{uploaded['id']}/split", json={"ranges": ["1"]})
    assert split.status_code == 200
    assert split.content.startswith(b"PK")
    history = client.get(f"/api/files/{uploaded['id']}/history")
    assert history.status_code == 200
    assert any(item["operation"] == "Sayfa döndürüldü" for item in history.json()["data"])


def test_image_conversion_preserves_transparency_with_background(client, sample_png):
    with sample_png.open("rb") as handle:
        uploaded = client.post("/api/files", files={"file": (sample_png.name, handle, "image/png")}).json()["data"]
    response = client.post(f"/api/images/{uploaded['id']}/convert", data={
        "output_format": "JPEG", "quality": "80", "background": "#ffffff"
    })
    assert response.status_code == 200, response.text
    assert response.content.startswith(b"\xff\xd8")

    with sample_png.open("rb") as handle:
        uploaded = client.post("/api/files", files={"file": (sample_png.name, handle, "image/png")}).json()["data"]
    webp = client.post(f"/api/images/{uploaded['id']}/convert", data={
        "output_format": "WEBP", "quality": "90",
    })
    assert webp.status_code == 200, webp.text
    from PIL import Image
    with Image.open(io.BytesIO(webp.content)) as converted:
        assert "A" in converted.getbands()
        assert 120 <= converted.getpixel((20, 20))[3] <= 136


def test_rejects_zero_page_pdf(client, tmp_path):
    from pypdf import PdfWriter
    empty = tmp_path / "empty.pdf"
    writer = PdfWriter()
    with empty.open("wb") as handle:
        writer.write(handle)
    with empty.open("rb") as handle:
        response = client.post("/api/files", files={"file": (empty.name, handle, "application/pdf")})
    assert response.status_code == 422
    assert "sayfa içermiyor" in response.json()["detail"]


def test_delete_and_reorder_pages_persist(client, multipage_pdf):
    uploaded = upload_pdf(client, multipage_pdf)
    reordered = client.post(f"/api/pdf/{uploaded['id']}/pages/reorder", json={"order": [2, 0, 1]})
    assert reordered.status_code == 200, reordered.text
    downloaded = client.get(f"/api/files/{uploaded['id']}/download").content
    with fitz.open(stream=downloaded, filetype="pdf") as output:
        assert "Sayfa 3" in output[0].get_text()

    deleted = client.post(f"/api/pdf/{uploaded['id']}/pages/delete", json={"pages": [1]})
    assert deleted.status_code == 200, deleted.text
    assert deleted.json()["data"]["page_count"] == 2


def test_blank_page_crop_and_page_decorations_persist(client, multipage_pdf):
    uploaded = upload_pdf(client, multipage_pdf)
    document_id = uploaded["id"]
    blank = client.post(f"/api/pdf/{document_id}/pages/blank", json={"after_page": 0})
    assert blank.status_code == 200, blank.text
    assert blank.json()["data"]["page_count"] == 4
    cropped = client.post(f"/api/pdf/{document_id}/pages/crop", json={"page": 0, "rect": [20, 20, 280, 360]})
    assert cropped.status_code == 200, cropped.text
    assert cropped.json()["data"]["pages"][0]["width"] == 260
    decorated = client.post(f"/api/pdf/{document_id}/pages/decorate", json={
        "header": "Atelier Baslik", "footer": "Gizli", "page_numbers": True, "start_number": 5,
    })
    assert decorated.status_code == 200, decorated.text
    content = client.get(f"/api/files/{document_id}/download").content
    with fitz.open(stream=content, filetype="pdf") as output:
        text = "\n".join(page.get_text().replace("\u00a0", " ") for page in output)
        assert "Atelier Baslik" in text and "Gizli" in text and "5" in text


def test_cannot_delete_every_page(client, sample_pdf):
    uploaded = upload_pdf(client, sample_pdf)
    response = client.post(f"/api/pdf/{uploaded['id']}/pages/delete", json={"pages": [0]})
    assert response.status_code == 400
    assert "en az bir sayfa" in response.json()["detail"]


def test_pdf_to_docx_is_editable_package(client, sample_pdf):
    uploaded = upload_pdf(client, sample_pdf)
    response = client.get(f"/api/convert/pdf/{uploaded['id']}/docx")
    assert response.status_code == 200, response.text
    assert zipfile.is_zipfile(io.BytesIO(response.content))
    from docx import Document
    document = Document(io.BytesIO(response.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Eski metin 12345" in text


def test_pdf_to_docx_keeps_table_cells_at_pdf_coordinates(client, table_layout_pdf):
    uploaded = upload_pdf(client, table_layout_pdf)
    response = client.get(f"/api/convert/pdf/{uploaded['id']}/docx")
    assert response.status_code == 200, response.text
    assert response.headers["X-Atelier-Mode"] == "editable-positioned-layout"

    from docx import Document
    from docx.oxml.ns import qn
    document = Document(io.BytesIO(response.content))
    positioned = {}
    for paragraph in document.paragraphs:
        frame = paragraph._p.pPr.find(qn("w:framePr")) if paragraph._p.pPr is not None else None
        if frame is not None and paragraph.text:
            positioned[paragraph.text] = (
                int(frame.get(qn("w:x"))) / 20,
                int(frame.get(qn("w:y"))) / 20,
            )

    with fitz.open(table_layout_pdf) as source:
        expected = {}
        for block in source[0].get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    expected[span["text"]] = span["bbox"][:2]

    for value in ["TOPLAM (TOTAL) (TL)", "3.142,15", "8.142,15", "ALT SATIR"]:
        assert value in positioned
        assert abs(positioned[value][0] - expected[value][0]) < 0.1
        assert abs(positioned[value][1] - expected[value][1]) < 0.1

    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert any(name.startswith("word/media/") for name in archive.namelist())


def test_pdf_to_docx_creates_one_positioned_section_per_page(client, multipage_pdf):
    uploaded = upload_pdf(client, multipage_pdf)
    response = client.get(f"/api/convert/pdf/{uploaded['id']}/docx")
    assert response.status_code == 200, response.text

    from docx import Document
    document = Document(io.BytesIO(response.content))
    assert len(document.sections) == 3
    assert {paragraph.text for paragraph in document.paragraphs if paragraph.text} >= {
        "Sayfa 1", "Sayfa 2", "Sayfa 3",
    }
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
        headers = [name for name in archive.namelist() if name.startswith("word/header") and name.endswith(".xml")]
        assert media
        assert len(headers) == 3


def test_pdf_to_images_creates_one_image_per_page(client, multipage_pdf):
    uploaded = upload_pdf(client, multipage_pdf)
    response = client.get(f"/api/convert/pdf/{uploaded['id']}/images?output_format=png&dpi=96")
    assert response.status_code == 200, response.text
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert archive.namelist() == ["sayfa-0001.png", "sayfa-0002.png", "sayfa-0003.png"]
        assert archive.read("sayfa-0001.png").startswith(b"\x89PNG")


def test_pdf_to_png_matches_opaque_source_render_pixel_for_pixel(client, table_layout_pdf):
    uploaded = upload_pdf(client, table_layout_pdf)
    response = client.get(f"/api/convert/pdf/{uploaded['id']}/images?output_format=png&dpi=144")
    assert response.status_code == 200, response.text
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        converted = archive.read("sayfa-0001.png")
    with fitz.open(table_layout_pdf) as source:
        expected = source[0].get_pixmap(matrix=fitz.Matrix(2, 2), colorspace=fitz.csRGB, alpha=False).tobytes("png")
    assert converted == expected
    from PIL import Image
    with Image.open(io.BytesIO(converted)) as image:
        assert image.mode == "RGB"


def test_images_to_pdf_and_docx_to_pdf(client, sample_png, sample_docx):
    with sample_png.open("rb") as first, sample_png.open("rb") as second:
        response = client.post("/api/convert/images/pdf", files=[
            ("files", ("bir.png", first, "image/png")),
            ("files", ("iki.png", second, "image/png")),
        ])
    assert response.status_code == 200, response.text
    with fitz.open(stream=response.content, filetype="pdf") as output:
        assert output.page_count == 2

    with sample_docx.open("rb") as handle:
        uploaded = client.post("/api/files", files={"file": (sample_docx.name, handle, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert uploaded.status_code == 200, uploaded.text
    document_id = uploaded.json()["data"]["id"]
    converted = client.get(f"/api/convert/docx/{document_id}/pdf")
    assert converted.status_code == 200, converted.text
    with fitz.open(stream=converted.content, filetype="pdf") as output:
        text = "\n".join(page.get_text() for page in output).replace("\u00a0", " ")
        assert "Atelier Docs" in text
        assert "Türkçe" in text


def test_office_files_can_be_inspected_edited_downloaded_and_exported(client, sample_docx, sample_xlsx, sample_pptx):
    fixtures = [(sample_docx, "document"), (sample_xlsx, "spreadsheet"), (sample_pptx, "presentation")]
    for path, expected_type in fixtures:
        with path.open("rb") as handle:
            uploaded = client.post("/api/files", files={"file": (path.name, handle, "application/octet-stream")})
        assert uploaded.status_code == 200, uploaded.text
        document_id = uploaded.json()["data"]["id"]
        inspected = client.get(f"/api/office/{document_id}")
        assert inspected.status_code == 200, inspected.text
        assert inspected.json()["data"]["type"] == expected_type
        if path.suffix == ".docx":
            updates = [{"kind": "paragraph", "index": 1, "value": "Güncellenmiş Türkçe paragraf"}]
        elif path.suffix == ".xlsx":
            updates = [{"sheet": "Veriler", "cell": "A2", "value": "Defter"}]
        else:
            updates = [{"slide": 0, "shape": 0, "value": "Güncel Sunum"}]
        saved = client.post(f"/api/office/{document_id}", json={"updates": updates})
        assert saved.status_code == 200, saved.text
        assert len(client.get(f"/api/files/{document_id}/download").content) > 100
        exported = client.get(f"/api/convert/office/{document_id}/pdf")
        assert exported.status_code == 200, exported.text
        assert exported.content.startswith(b"%PDF")


def test_office_structural_edits_persist(client, sample_docx, sample_pptx):
    with sample_docx.open("rb") as handle:
        uploaded = client.post("/api/files", files={"file": (sample_docx.name, handle, "application/octet-stream")})
    document_id = uploaded.json()["data"]["id"]
    saved = client.post(f"/api/office/{document_id}", json={"updates": [
        {"kind": "add_paragraph", "value": "Ek paragraf"},
        {"kind": "add_table", "rows": 2, "columns": 3},
    ]})
    assert saved.status_code == 200, saved.text
    inspected = client.get(f"/api/office/{document_id}").json()["data"]
    assert inspected["paragraphs"][-1]["text"] == "Ek paragraf"
    assert len(inspected["tables"][-1]["rows"]) == 2
    assert len(inspected["tables"][-1]["rows"][0]) == 3

    with sample_pptx.open("rb") as handle:
        uploaded = client.post("/api/files", files={"file": (sample_pptx.name, handle, "application/octet-stream")})
    presentation_id = uploaded.json()["data"]["id"]
    added = client.post(f"/api/office/{presentation_id}", json={"updates": [
        {"kind": "add_slide", "value": "İkinci Slayt"},
    ]})
    assert added.status_code == 200, added.text
    inspected = client.get(f"/api/office/{presentation_id}").json()["data"]
    assert len(inspected["slides"]) == 2
    assert inspected["slides"][1]["texts"][0]["text"] == "İkinci Slayt"
    reordered = client.post(f"/api/office/{presentation_id}", json={"updates": [
        {"kind": "reorder_slides", "order": [1, 0]},
    ]})
    assert reordered.status_code == 200, reordered.text
    inspected = client.get(f"/api/office/{presentation_id}").json()["data"]
    assert inspected["slides"][0]["texts"][0]["text"] == "İkinci Slayt"


def test_pdf_exports_editable_xlsx_and_pptx(client, sample_pdf):
    uploaded = upload_pdf(client, sample_pdf)
    document_id = uploaded["id"]
    xlsx = client.get(f"/api/convert/pdf/{document_id}/xlsx")
    pptx = client.get(f"/api/convert/pdf/{document_id}/pptx")
    assert xlsx.status_code == 200, xlsx.text
    assert pptx.status_code == 200, pptx.text
    from openpyxl import load_workbook
    from pptx import Presentation
    workbook = load_workbook(io.BytesIO(xlsx.content))
    editable_values = [cell.value for row in workbook["Sayfa 1"].iter_rows() for cell in row if cell.value]
    assert "Eski metin 12345" in editable_values
    assert workbook["Sayfa 1"]["A1"].value is None
    assert len(workbook["Görünüm 1"]._images) == 1
    presentation = Presentation(io.BytesIO(pptx.content))
    assert presentation.slides[0].shapes[0].shape_type == 13
    assert any("Eski metin 12345" in shape.text for shape in presentation.slides[0].shapes if hasattr(shape, "text"))


def test_pdf_to_pptx_keeps_distant_table_text_at_source_coordinates(client, table_layout_pdf):
    uploaded = upload_pdf(client, table_layout_pdf)
    response = client.get(f"/api/convert/pdf/{uploaded['id']}/pptx")
    assert response.status_code == 200, response.text
    assert response.headers["X-Atelier-Mode"] == "editable-positioned-layout"
    from pptx import Presentation
    presentation = Presentation(io.BytesIO(response.content))
    positions = {
        shape.text: (shape.left / 12700, shape.top / 12700)
        for shape in presentation.slides[0].shapes
        if getattr(shape, "has_text_frame", False) and shape.text
    }
    with fitz.open(table_layout_pdf) as source:
        expected = {
            span["text"]: span["bbox"][:2]
            for block in source[0].get_text("dict")["blocks"]
            for line in block.get("lines", [])
            for span in line.get("spans", [])
            if span.get("text", "").strip()
        }
    for value in ["TOPLAM (TOTAL) (TL)", "3.142,15", "8.142,15", "ALT SATIR"]:
        assert value in positions
        assert abs(positions[value][0] - expected[value][0]) < 0.1
        assert abs(positions[value][1] - expected[value][1]) < 0.1


def test_add_text_shape_and_permanent_redaction(client, sample_pdf):
    uploaded = upload_pdf(client, sample_pdf)
    document_id = uploaded["id"]

    added = client.post(f"/api/pdf/{document_id}/content/text", json={
        "page": 0,
        "rect": [70, 180, 360, 230],
        "text": "Eklenen Türkçe metin",
        "style": {"font_size": 12, "color": "#6d5dfc", "align": "left"},
    })
    assert added.status_code == 200, added.text

    shaped = client.post(f"/api/pdf/{document_id}/content/shape", json={
        "page": 0,
        "kind": "rectangle",
        "rect": [65, 170, 370, 240],
        "stroke_color": "#ef4444",
        "fill_color": None,
        "width": 2,
    })
    assert shaped.status_code == 200, shaped.text

    blocks = client.get(f"/api/pdf/{document_id}/pages/0/text").json()["data"]
    old = next(item for item in blocks if "Eski metin 12345" in item["text"])
    redacted = client.post(f"/api/pdf/{document_id}/content/redact", json={
        "page": 0,
        "rects": [old["bbox"]],
    })
    assert redacted.status_code == 200, redacted.text
    assert redacted.json()["data"]["audit"]["areas"] == 1

    downloaded = client.get(f"/api/files/{document_id}/download").content
    with fitz.open(stream=downloaded, filetype="pdf") as output:
        text = "\n".join(page.get_text() for page in output).replace("\u00a0", " ")
        assert "Eski metin 12345" not in text
        assert "Eklenen Türkçe metin" in text


def test_undo_and_redo_restore_pdf_revisions(client, sample_pdf):
    uploaded = upload_pdf(client, sample_pdf)
    document_id = uploaded["id"]
    payload = {
        "page": 0, "rect": [70, 180, 360, 230], "text": "Geri alınabilir metin",
        "style": {"font_size": 12, "color": "#111827", "align": "left"},
    }
    assert client.post(f"/api/pdf/{document_id}/content/text", json=payload).status_code == 200

    undone = client.post(f"/api/pdf/{document_id}/undo")
    assert undone.status_code == 200, undone.text
    content = client.get(f"/api/files/{document_id}/download").content
    with fitz.open(stream=content, filetype="pdf") as output:
        assert "Geri alınabilir metin" not in output[0].get_text()

    redone = client.post(f"/api/pdf/{document_id}/redo")
    assert redone.status_code == 200, redone.text
    content = client.get(f"/api/files/{document_id}/download").content
    with fitz.open(stream=content, filetype="pdf") as output:
        assert "Geri alınabilir metin" in output[0].get_text().replace("\u00a0", " ")


def test_add_image_embeds_real_pdf_image_object(client, sample_pdf, sample_png):
    uploaded = upload_pdf(client, sample_pdf)
    with sample_png.open("rb") as handle:
        response = client.post(f"/api/pdf/{uploaded['id']}/content/image", data={
            "page": "0", "rect": "[70, 250, 260, 380]",
        }, files={"file": (sample_png.name, handle, "image/png")})
    assert response.status_code == 200, response.text
    content = client.get(f"/api/files/{uploaded['id']}/download").content
    with fitz.open(stream=content, filetype="pdf") as output:
        assert len(output[0].get_images(full=True)) >= 1


def test_embedded_image_can_be_drag_moved_and_undone(client, sample_pdf, sample_png):
    uploaded = upload_pdf(client, sample_pdf)
    document_id = uploaded["id"]
    with sample_png.open("rb") as handle:
        inserted = client.post(f"/api/pdf/{document_id}/content/image", data={
            "page": "0", "rect": "[100, 180, 300, 330]",
        }, files={"file": (sample_png.name, handle, "image/png")})
    assert inserted.status_code == 200, inserted.text
    listed = client.get(f"/api/pdf/{document_id}/pages/0/images")
    assert listed.status_code == 200, listed.text
    image = listed.json()["data"][0]
    source = image["bbox"]
    width, height = source[2] - source[0], source[3] - source[1]
    target = [40.0, 350.0, 40.0 + width, 350.0 + height]

    moved = client.post(f"/api/pdf/{document_id}/content/image/move", json={
        "page": 0, "xref": image["xref"], "source_rect": source, "target_rect": target,
    })
    assert moved.status_code == 200, moved.text
    after = client.get(f"/api/pdf/{document_id}/pages/0/images").json()["data"]
    assert any(all(abs(a - b) < 1 for a, b in zip(item["bbox"], target)) for item in after)
    assert not any(all(abs(a - b) < 1 for a, b in zip(item["bbox"], source)) for item in after)

    undone = client.post(f"/api/pdf/{document_id}/undo")
    assert undone.status_code == 200, undone.text
    restored = client.get(f"/api/pdf/{document_id}/pages/0/images").json()["data"]
    assert any(all(abs(a - b) < 1 for a, b in zip(item["bbox"], source)) for item in restored)


def test_aes_encryption_unlock_and_metadata(client, sample_pdf):
    uploaded = upload_pdf(client, sample_pdf)
    protected = client.post(f"/api/pdf/{uploaded['id']}/security/encrypt", json={
        "user_password": "guclu-parola",
        "owner_password": "sahip-parolasi",
        "allow_print": True,
        "allow_copy": False,
        "allow_annotate": False,
    })
    assert protected.status_code == 200, protected.text
    with fitz.open(stream=protected.content, filetype="pdf") as encrypted:
        assert encrypted.needs_pass
        assert encrypted.authenticate("guclu-parola") > 0

    encrypted_file = io.BytesIO(protected.content)
    encrypted_upload = client.post("/api/files", files={"file": ("korumali.pdf", encrypted_file, "application/pdf")})
    assert encrypted_upload.status_code == 200, encrypted_upload.text
    encrypted_id = encrypted_upload.json()["data"]["id"]
    assert encrypted_upload.json()["data"]["pdf"]["encrypted"] is True
    assert client.post(f"/api/pdf/{encrypted_id}/security/unlock", json={"password": "yanlis"}).status_code == 400
    unlocked = client.post(f"/api/pdf/{encrypted_id}/security/unlock", json={"password": "guclu-parola"})
    assert unlocked.status_code == 200, unlocked.text
    assert unlocked.json()["data"]["encrypted"] is False

    metadata = client.post(f"/api/pdf/{encrypted_id}/metadata", json={
        "title": "Atelier Test Belgesi", "author": "QA", "subject": "Türkçe", "keywords": "pdf,test",
    })
    assert metadata.status_code == 200, metadata.text
    downloaded = client.get(f"/api/files/{encrypted_id}/download").content
    with fitz.open(stream=downloaded, filetype="pdf") as output:
        assert output.metadata["title"] == "Atelier Test Belgesi"
        assert output.metadata["author"] == "QA"


def test_balanced_compression_recompresses_embedded_images(client, image_heavy_pdf):
    uploaded = upload_pdf(client, image_heavy_pdf)
    response = client.post(f"/api/pdf/{uploaded['id']}/compress", json={"level": "balanced"})
    assert response.status_code == 200, response.text
    metrics = response.json()["data"]
    assert metrics["optimized_images"] >= 1
    assert metrics["output_size"] < metrics["original_size"]
    content = client.get(f"/api/files/{uploaded['id']}/download").content
    with fitz.open(stream=content, filetype="pdf") as output:
        assert output.page_count == 1
        assert len(output[0].get_images()) >= 1


def test_pdf_compression_job_reports_progress_and_result(client, image_heavy_pdf):
    import time
    uploaded = upload_pdf(client, image_heavy_pdf)
    queued = client.post(f"/api/jobs/pdf/{uploaded['id']}/compress", json={"level": "balanced"})
    assert queued.status_code == 200, queued.text
    job_id = queued.json()["data"]["id"]
    job = queued.json()["data"]
    for _ in range(100):
        job = client.get(f"/api/jobs/{job_id}").json()["data"]
        if job["status"] in {"completed", "failed", "cancelled"}:
            break
        time.sleep(.03)
    assert job["status"] == "completed", job
    assert job["progress"] == 100
    assert job["result"]["output_size"] <= job["result"]["original_size"]


def test_image_target_size_compression(client, tmp_path):
    from PIL import Image
    image_path = tmp_path / "buyuk.png"
    Image.effect_noise((1200, 900), 100).convert("RGB").save(image_path, format="PNG")
    with image_path.open("rb") as handle:
        uploaded = client.post("/api/files", files={"file": (image_path.name, handle, "image/png")}).json()["data"]
    response = client.post(f"/api/images/{uploaded['id']}/compress", data={
        "target_kb": "100", "output_format": "JPEG",
    })
    assert response.status_code == 200, response.text
    assert response.content.startswith(b"\xff\xd8")
    assert len(response.content) <= 100 * 1024


def test_batch_image_conversion_and_pdf_compression_return_zip(client, sample_png, sample_pdf):
    with sample_png.open("rb") as first, sample_png.open("rb") as second:
        images = client.post("/api/batch/images/convert", data={"output_format": "WEBP", "quality": "75"}, files=[
            ("files", ("bir.png", first, "image/png")), ("files", ("iki.png", second, "image/png")),
        ])
    assert images.status_code == 200, images.text
    with zipfile.ZipFile(io.BytesIO(images.content)) as archive:
        assert sorted(archive.namelist()) == ["bir.webp", "iki.webp"]
    with sample_pdf.open("rb") as first, sample_pdf.open("rb") as second:
        pdfs = client.post("/api/batch/pdf/compress", data={"level": "quality"}, files=[
            ("files", ("bir.pdf", first, "application/pdf")), ("files", ("iki.pdf", second, "application/pdf")),
        ])
    assert pdfs.status_code == 200, pdfs.text
    with zipfile.ZipFile(io.BytesIO(pdfs.content)) as archive:
        assert len(archive.namelist()) == 2


def test_image_crop_and_filters_are_written_to_pixels(client, sample_png):
    with sample_png.open("rb") as handle:
        uploaded = client.post("/api/files", files={"file": (sample_png.name, handle, "image/png")}).json()["data"]
    response = client.post(f"/api/images/{uploaded['id']}/edit", data={
        "crop_x": "10", "crop_y": "5", "crop_width": "60", "crop_height": "40",
        "grayscale": "false", "sepia": "true", "sharpen": "1.5", "blur": "0.5",
    })
    assert response.status_code == 200, response.text
    from PIL import Image
    with Image.open(io.BytesIO(response.content)) as output:
        assert output.size == (60, 40)
        pixel = output.convert("RGB").getpixel((20, 20))
        assert pixel[0] >= pixel[2]


def test_freehand_ink_is_written_as_pdf_vector_content(client, sample_pdf):
    uploaded = upload_pdf(client, sample_pdf)
    response = client.post(f"/api/pdf/{uploaded['id']}/content/ink", json={
        "page": 0, "paths": [[[80, 220], [100, 235], [125, 218], [155, 242]]],
        "color": "#2345aa", "width": 3,
    })
    assert response.status_code == 200, response.text
    content = client.get(f"/api/files/{uploaded['id']}/download").content
    with fitz.open(stream=content, filetype="pdf") as output:
        assert any(drawing.get("color") is not None for drawing in output[0].get_drawings())


def test_ocr_reports_missing_local_engine_without_modifying_document(client, sample_pdf, monkeypatch):
    import backend.app.services.ocr_service as ocr_module
    monkeypatch.setattr(ocr_module, "find_tesseract", lambda: None)
    uploaded = upload_pdf(client, sample_pdf)
    response = client.post(f"/api/pdf/{uploaded['id']}/ocr", json={"language": "tur+eng", "dpi": 150})
    assert response.status_code == 503
    assert "OCR motoru" in response.json()["detail"]
    content = client.get(f"/api/files/{uploaded['id']}/download").content
    with fitz.open(stream=content, filetype="pdf") as output:
        assert "Eski metin 12345" in output[0].get_text()


def test_scanned_image_becomes_searchable_pdf(client, scanned_document_png):
    with scanned_document_png.open("rb") as handle:
        response = client.post(
            "/api/ocr/images/pdf",
            files=[("files", (scanned_document_png.name, handle, "image/png"))],
            data={"language": "tur+eng", "dpi": "200"},
        )
    assert response.status_code == 200, response.text
    assert response.content.startswith(b"%PDF")
    with fitz.open(stream=response.content, filetype="pdf") as output:
        assert output.page_count == 1
        text = output[0].get_text().casefold()
        assert "12345" in text
        assert "atelier" in text
        assert output[0].get_images()


def test_scanned_pdf_is_detected_and_ocr_layer_persists(client, scanned_document_png, tmp_path):
    scanned_pdf = tmp_path / "tarama.pdf"
    document = fitz.open()
    page = document.new_page(width=504, height=306)
    page.insert_image(page.rect, filename=str(scanned_document_png))
    document.save(scanned_pdf)
    document.close()
    uploaded = upload_pdf(client, scanned_pdf)
    assert uploaded["pdf"]["scanned_likely"] is True
    assert uploaded["pdf"]["pages"][0]["scanned_likely"] is True
    response = client.post(f"/api/pdf/{uploaded['id']}/ocr", json={"language": "tur+eng", "dpi": 200})
    assert response.status_code == 200, response.text
    assert response.json()["data"]["audit"]["pages_processed"] == 1
    content = client.get(f"/api/files/{uploaded['id']}/download").content
    with fitz.open(stream=content, filetype="pdf") as output:
        assert "12345" in output[0].get_text()


def test_removable_watermark_and_clickable_link(client, sample_pdf):
    uploaded = upload_pdf(client, sample_pdf)
    document_id = uploaded["id"]
    watermark = client.post(f"/api/pdf/{document_id}/content/watermark", json={
        "text": "GIZLI", "font_size": 36, "color": "#999999", "opacity": 0.3, "rotation": 0,
    })
    assert watermark.status_code == 200, watermark.text
    link = client.post(f"/api/pdf/{document_id}/content/link", json={
        "page": 0, "rect": [70, 160, 260, 190], "uri": "https://example.com",
    })
    assert link.status_code == 200, link.text
    content = client.get(f"/api/files/{document_id}/download").content
    with fitz.open(stream=content, filetype="pdf") as output:
        assert any((annot.info or {}).get("subject") == "AtelierWatermark" for annot in (output[0].annots() or []))
        assert any(item.get("uri") == "https://example.com" for item in output[0].get_links())
    removed = client.post(f"/api/pdf/{document_id}/content/watermark/remove", json={})
    assert removed.status_code == 200, removed.text
    content = client.get(f"/api/files/{document_id}/download").content
    with fitz.open(stream=content, filetype="pdf") as output:
        assert not any((annot.info or {}).get("subject") == "AtelierWatermark" for annot in (output[0].annots() or []))
