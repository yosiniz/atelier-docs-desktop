import fitz
import pytest
from fastapi.testclient import TestClient

from backend.app.main import app


@pytest.fixture()
def client() -> TestClient:
    return TestClient(app)


@pytest.fixture()
def sample_pdf(tmp_path):
    path = tmp_path / "turkce-belge.pdf"
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.insert_text((72, 96), "Eski metin 12345", fontsize=14)
    page.insert_text((72, 140), "Ikinci satir", fontsize=12)
    document.save(path)
    document.close()
    return path


@pytest.fixture()
def table_layout_pdf(tmp_path):
    """A table-like page where PyMuPDF groups distant cells in one block."""
    path = tmp_path / "tablo-yerlesimi.pdf"
    document = fitz.open()
    page = document.new_page(width=520, height=200)
    page.insert_text((280, 80), "TOPLAM (TOTAL) (TL)", fontsize=10, fontname="hebo")
    page.insert_text((440, 80), "3.142,15", fontsize=10)
    page.insert_text((40, 120), "8.142,15", fontsize=10)
    page.insert_text((40, 132), "ALT SATIR", fontsize=10, fontname="hebo")
    document.save(path)
    document.close()
    return path


@pytest.fixture()
def sample_png(tmp_path):
    from PIL import Image
    path = tmp_path / "ornek.png"
    Image.new("RGBA", (120, 80), (255, 0, 80, 128)).save(path)
    return path


@pytest.fixture()
def scanned_document_png(tmp_path):
    from PIL import Image, ImageDraw, ImageFont
    path = tmp_path / "taranmis-belge.png"
    image = Image.new("RGB", (1400, 850), "white")
    draw = ImageDraw.Draw(image)
    font_path = "C:/Windows/Fonts/arial.ttf"
    try:
        heading = ImageFont.truetype(font_path, 62)
        body = ImageFont.truetype(font_path, 42)
    except OSError:
        heading = body = ImageFont.load_default()
    draw.text((100, 110), "TARANMIS BELGE", fill="black", font=heading)
    draw.text((100, 240), "Fatura numarasi 12345", fill="black", font=body)
    draw.text((100, 320), "Atelier Docs OCR testi", fill="black", font=body)
    image.save(path, dpi=(200, 200))
    return path


@pytest.fixture()
def multipage_pdf(tmp_path):
    path = tmp_path / "cok-sayfali.pdf"
    document = fitz.open()
    for number in range(1, 4):
        page = document.new_page(width=300, height=400)
        page.insert_text((40, 70), f"Sayfa {number}", fontsize=18)
    document.save(path)
    document.close()
    return path


@pytest.fixture()
def sample_docx(tmp_path):
    from docx import Document
    path = tmp_path / "ornek.docx"
    document = Document()
    document.add_heading("Atelier Docs Başlık", level=1)
    document.add_paragraph("Türkçe içerik: ç, ğ, ı, İ, ö, ş, ü")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Ürün"
    table.cell(0, 1).text = "Adet"
    table.cell(1, 0).text = "Kalem"
    table.cell(1, 1).text = "3"
    document.save(path)
    return path


@pytest.fixture()
def sample_xlsx(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "ornek.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Veriler"
    sheet["A1"], sheet["B1"] = "Ürün", "Adet"
    sheet["A2"], sheet["B2"] = "Kalem", 3
    sheet["B3"] = "=SUM(B2:B2)"
    workbook.save(path)
    return path


@pytest.fixture()
def sample_pptx(tmp_path):
    from pptx import Presentation
    path = tmp_path / "ornek.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Atelier Sunum"
    slide.placeholders[1].text = "Düzenlenebilir içerik"
    presentation.save(path)
    return path


@pytest.fixture()
def image_heavy_pdf(tmp_path):
    from PIL import Image
    image_path = tmp_path / "noise.png"
    image = Image.effect_noise((1200, 1200), 100).convert("RGB")
    image.save(image_path, format="PNG")
    path = tmp_path / "gorselli.pdf"
    document = fitz.open()
    page = document.new_page(width=600, height=600)
    page.insert_image(page.rect, filename=str(image_path))
    document.save(path)
    document.close()
    return path
