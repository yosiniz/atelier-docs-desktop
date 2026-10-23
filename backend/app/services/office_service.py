from __future__ import annotations

import csv
from pathlib import Path

import fitz
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

from .native_office_service import office_to_pdf_native
from .pdf_service import _font_name, _unicode_font


def inspect_office(path: Path) -> dict:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        document = Document(path)
        return {
            "type": "document",
            "paragraphs": [
                {"index": index, "text": paragraph.text, "style": paragraph.style.name if paragraph.style else "Normal"}
                for index, paragraph in enumerate(document.paragraphs)
            ],
            "tables": [
                {"index": table_index, "rows": [[cell.text for cell in row.cells] for row in table.rows]}
                for table_index, table in enumerate(document.tables)
            ],
        }
    if suffix in {".xlsx", ".csv"}:
        if suffix == ".csv":
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle))[:500]
            return {"type": "spreadsheet", "sheets": [{"name": "CSV", "rows": rows, "truncated": len(rows) == 500}]}
        workbook = load_workbook(path, data_only=False, read_only=False)
        sheets = []
        for sheet in workbook.worksheets:
            max_row = min(sheet.max_row, 500)
            max_col = min(sheet.max_column, 50)
            rows = [[sheet.cell(row, column).value for column in range(1, max_col + 1)] for row in range(1, max_row + 1)]
            sheets.append({"name": sheet.title, "rows": rows, "truncated": sheet.max_row > 500 or sheet.max_column > 50})
        workbook.close()
        return {"type": "spreadsheet", "sheets": sheets}
    if suffix == ".pptx":
        presentation = Presentation(path)
        slides = []
        for slide_index, slide in enumerate(presentation.slides):
            items = []
            for shape_index, shape in enumerate(slide.shapes):
                if getattr(shape, "has_text_frame", False):
                    items.append({"shape": shape_index, "text": shape.text, "name": shape.name})
            slides.append({"index": slide_index, "texts": items})
        return {"type": "presentation", "slides": slides}
    raise ValueError("Desteklenmeyen Office dosyası")


def save_office(path: Path, output: Path, updates: list[dict]) -> None:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        document = Document(path)
        for update in updates:
            kind = update.get("kind")
            value = str(update.get("value", ""))
            if kind == "paragraph":
                document.paragraphs[int(update["index"])].text = value
            elif kind == "table":
                document.tables[int(update["table"])].cell(int(update["row"]), int(update["column"])).text = value
            elif kind == "add_paragraph":
                document.add_paragraph(value)
            elif kind == "add_table":
                rows, columns = max(1, int(update.get("rows", 2))), max(1, int(update.get("columns", 2)))
                document.add_table(rows=rows, cols=columns)
        document.save(output)
        return
    if suffix == ".xlsx":
        workbook = load_workbook(path, data_only=False)
        for update in updates:
            sheet = workbook[str(update["sheet"])]
            sheet[str(update["cell"])].value = update.get("value", "")
        workbook.save(output)
        workbook.close()
        return
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        for update in updates:
            row, column = int(update["row"]), int(update["column"])
            while len(rows) <= row:
                rows.append([])
            while len(rows[row]) <= column:
                rows[row].append("")
            rows[row][column] = str(update.get("value", ""))
        with output.open("w", encoding="utf-8-sig", newline="") as handle:
            csv.writer(handle).writerows(rows)
        return
    if suffix == ".pptx":
        presentation = Presentation(path)
        for update in updates:
            kind = update.get("kind", "text")
            if kind == "add_slide":
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                box = slide.shapes.add_textbox(Inches(.7), Inches(.6), presentation.slide_width - Inches(1.4), Inches(1.0))
                box.text_frame.text = str(update.get("value", "Yeni Slayt"))
                box.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
                continue
            if kind == "delete_slide":
                if len(presentation.slides) <= 1:
                    raise ValueError("Sunumda en az bir slayt kalmalıdır")
                slide_id = presentation.slides._sldIdLst[int(update["slide"])]
                presentation.part.drop_rel(slide_id.rId)
                presentation.slides._sldIdLst.remove(slide_id)
                continue
            if kind in {"reorder", "reorder_slides"}:
                order = [int(index) for index in update["order"]]
                if sorted(order) != list(range(len(presentation.slides))):
                    raise ValueError("Slayt sırası bütün slaytları içermelidir")
                slide_ids = list(presentation.slides._sldIdLst)
                for slide_id in list(presentation.slides._sldIdLst):
                    presentation.slides._sldIdLst.remove(slide_id)
                for index in order:
                    presentation.slides._sldIdLst.append(slide_ids[index])
                continue
            slide = presentation.slides[int(update["slide"])]
            shape = slide.shapes[int(update["shape"])]
            if not getattr(shape, "has_text_frame", False):
                raise ValueError("Seçilen PowerPoint nesnesi metin içermiyor")
            shape.text = str(update.get("value", ""))
        presentation.save(output)
        return
    raise ValueError("Desteklenmeyen Office dosyası")


def _new_writer_page(pdf: fitz.Document, width: float = 595, height: float = 842):
    page = pdf.new_page(width=width, height=height)
    return page, 50.0


def spreadsheet_to_pdf(path: Path, output: Path) -> dict:
    native = office_to_pdf_native(path, output)
    if native is not None:
        return native

    data = inspect_office(path)
    pdf = fitz.open()
    font_path = _unicode_font()
    font_name = _font_name() if font_path else "helv"
    for sheet in data["sheets"]:
        page, y = _new_writer_page(pdf, 842, 595)
        page.insert_text((36, y), sheet["name"], fontsize=16, fontname=font_name,
                         fontfile=str(font_path) if font_path else None, color=(0.12, 0.14, 0.22))
        y += 26
        for row in sheet["rows"]:
            text = " | ".join("" if value is None else str(value) for value in row)
            if y > page.rect.height - 35:
                page, y = _new_writer_page(pdf, 842, 595)
            page.insert_textbox(fitz.Rect(36, y, page.rect.width - 36, y + 18), text, fontsize=8,
                                fontname=font_name, fontfile=str(font_path) if font_path else None,
                                color=(0.15, 0.17, 0.23))
            y += 15
    pdf.save(output, garbage=4, deflate=True)
    pdf.close()
    return {"mode": "local-spreadsheet-render"}


def presentation_to_pdf(path: Path, output: Path) -> dict:
    native = office_to_pdf_native(path, output)
    if native is not None:
        return native

    presentation = Presentation(path)
    pdf = fitz.open()
    width = presentation.slide_width / 12700
    height = presentation.slide_height / 12700
    font_path = _unicode_font()
    font_name = _font_name() if font_path else "helv"
    for slide in presentation.slides:
        page = pdf.new_page(width=width, height=height)
        page.draw_rect(page.rect, fill=(1, 1, 1), color=None)
        for shape in slide.shapes:
            rect = fitz.Rect(shape.left / 12700, shape.top / 12700,
                             (shape.left + shape.width) / 12700, (shape.top + shape.height) / 12700)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    page.insert_image(rect, stream=shape.image.blob, keep_proportion=True)
                except Exception:
                    pass
            elif getattr(shape, "has_text_frame", False) and shape.text.strip():
                size = 16.0
                for paragraph in shape.text_frame.paragraphs:
                    if paragraph.runs and paragraph.runs[0].font.size:
                        size = max(6, min(48, paragraph.runs[0].font.size.pt))
                        break
                page.insert_textbox(rect, shape.text, fontsize=size, fontname=font_name,
                                    fontfile=str(font_path) if font_path else None,
                                    color=(0.08, 0.1, 0.16), overlay=True)
    pdf.save(output, garbage=4, deflate=True)
    pdf.close()
    return {"mode": "local-presentation-render"}
